#!/usr/bin/env python
"""Build the 6-slide LipiLens presentation (heritage theme, 16:9).

Usage:  python presentation/build_ppt.py
Output: presentation/lipilens_presentation.pptx

All numbers are the project's real measured results. Author placeholders
marked TBD — replace before presenting.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentation" / "lipilens_presentation.pptx"

CREAM = RGBColor(0xF6, 0xF1, 0xE7)
INK = RGBColor(0x2B, 0x26, 0x20)
GOLD = RGBColor(0xA8, 0x7F, 0x2E)
GREEN = RGBColor(0x1E, 0x3A, 0x2F)
MUTED = RGBColor(0x7A, 0x72, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB3, 0x26, 0x1E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, l, t, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = WHITE
    sp.line.fill.background()
    return sp


def text(slide, l, t, w, h, runs, size=20, bold=False, color=INK,
         align=PP_ALIGN.LEFT, font="Georgia"):
    tx = slide.shapes.add_textbox(l, t, w, h).text_frame
    tx.word_wrap = True
    p = tx.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, {})]
    for s, fmt in runs:
        r = p.add_run()
        r.text = s
        r.font.size = Pt(fmt.get("size", size))
        r.font.bold = fmt.get("bold", bold)
        r.font.color.rgb = fmt.get("color", color)
        r.font.name = fmt.get("font", font)
    return tx


def bullets(slide, l, t, w, h, items, size=18):
    tx = slide.shapes.add_textbox(l, t, w, h).text_frame
    tx.word_wrap = True
    for i, it in enumerate(items):
        p = tx.paragraphs[0] if i == 0 else tx.add_paragraph()
        p.space_after = Pt(8)
        p.level = 0
        if isinstance(it, str):
            it = [(it, {})]
        norm = [(el if isinstance(el, tuple) else (el, {})) for el in it]
        for j, (s, fmt) in enumerate(norm):
            r = p.add_run()
            r.text = ("•  " if j == 0 else "") + s
            r.font.size = Pt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", False)
            r.font.color.rgb = fmt.get("color", INK)
            r.font.name = "Georgia"
    return tx


def title_bar(slide, kicker, title):
    text(slide, Inches(0.7), Inches(0.25), Inches(11.9), Inches(0.5),
         kicker, size=15, color=GOLD, align=PP_ALIGN.LEFT)
    text(slide, Inches(0.7), Inches(0.65), Inches(11.9), Inches(0.9),
         title, size=34, bold=True, align=PP_ALIGN.LEFT)


def footer(slide, n):
    text(slide, Inches(0.7), Inches(6.9), Inches(6), Inches(0.4),
         "LipiLens · Kolhapur Institute of Technology, Kolhapur",
         size=12, color=MUTED)
    text(slide, Inches(11.6), Inches(6.9), Inches(1), Inches(0.4), str(n),
         size=12, color=MUTED, align=PP_ALIGN.RIGHT)


# ---- 1. Title -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
text(s, Inches(1), Inches(1.4), Inches(11.3), Inches(1.2), "LipiLens",
     size=72, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(2.5), Inches(11.3), Inches(0.7),
     "Ancient Scripts. New Possibilities.", size=28, color=GOLD,
     align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(3.5), Inches(11.3), Inches(0.9),
     "AI-assisted transcription of historical Modi manuscripts — with a "
     "measured answer to whether image restoration actually helps.",
     size=20, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(5.0), Inches(11.3), Inches(0.5), "Chirayu Khot",
     size=22, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(5.5), Inches(11.3), Inches(0.5),
     "Kolhapur Institute of Technology, Kolhapur  ·  (+ 3 co-authors & "
     "faculty mentor — TBD)", size=15, color=MUTED, align=PP_ALIGN.CENTER)

# ---- 2. Problem -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
title_bar(s, "01 · PROBLEM", "Millions of pages no machine can read")
bullets(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(4.5), [
    [("Modi script: ", {"bold": True}),
     ("Maharashtra's administrative script, 13th c. to mid-20th c. — land, "
      "court and personal records survive only in Modi.")],
    [("Readers are disappearing, pages are degrading: ", {"bold": True}),
     ("faded ink, stains, skew, bad photography.")],
    [("No mature OCR exists for Modi. ", {"bold": True}),
     ("Vision-language models + LoRA adapters can draft transliterations — "
      "but a fluent draft can be entirely wrong.")],
    [("Two questions: ", {"bold": True}),
     ("can we build a trustworthy human-in-the-loop pipeline, and does "
      "classical restoration help or hurt the model?")],
])
footer(s, 2)

# ---- 3. System ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
title_bar(s, "02 · SYSTEM", "LipiLens: upload → restore → transcribe → verify")
steps = ["Upload\n(photo)", "Restore\n(OpenCV, 7 configs)", "Transcribe\n(Qwen2.5-VL-3B + Modi LoRA)",
         "Verify\n(human, explicit)", "Archive\n+ search"]
for i, st in enumerate(steps):
    x = Inches(0.7 + i * 2.5)
    b = box(s, x, Inches(2.0), Inches(2.1), Inches(1.6))
    b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = b.text_frame.paragraphs[0].add_run()
    r.text = st
    r.font.size = Pt(16)
    r.font.name = "Georgia"
    r.font.color.rgb = INK
    if i < 4:
        text(s, x + Inches(2.1), Inches(2.55), Inches(0.4), Inches(0.5), "→",
             size=28, color=GOLD, align=PP_ALIGN.CENTER)
bullets(s, Inches(0.7), Inches(4.2), Inches(11.9), Inches(2.3), [
    [("AI-authoritative use is structurally impossible: ", {"bold": True}),
     ("the AI draft is immutable; verification writes a separate field.")],
    [("Originals are never overwritten; ", {"bold": True}),
     ("every restoration config is versioned for reproducibility.")],
    "Live demo: upload → compare slider → edit → verify → library.",
])
footer(s, 3)

# ---- 4. Experiment --------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
title_bar(s, "03 · EXPERIMENT", "140 calls, 7 conditions, 20 pages")
bullets(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(4.5), [
    [("Data: ", {"bold": True}),
     ("50 real MoDeTrans pages (IIT Roorkee, MIT) with expert Devanagari "
      "ground truth; sweep on 20, refs 60–205 chars.")],
    [("Model: ", {"bold": True}),
     ("Qwen2.5-VL-3B + Modi LoRA (pinned revisions), official prompt, "
      "greedy decoding, Colab GPU.")],
    [("Conditions: ", {"bold": True}),
     ("original (control) · grayscale · denoised · enhanced (CLAHE) · "
      "binarized (Otsu) · deskewed · full restoration.")],
    [("Metrics: ", {"bold": True}),
     ("CER / WER vs ground truth; same model+prompt+decoding, only the "
      "input varies. Baseline n=50: CER 0.304.")],
])
footer(s, 4)

# ---- 5. Results -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
title_bar(s, "04 · RESULTS", "On clean pages, send the original")
s.shapes.add_picture(
    str(ROOT / "docs" / "sweep_means.png"),
    Inches(0.7), Inches(1.8), width=Inches(6.6))
bullets(s, Inches(7.7), Inches(1.8), Inches(4.9), Inches(4.8), [
    [("Grayscale = perfect no-op ", {"bold": True}), ("(20/20 ties: the control that validates the rig).")],
    [("Binarization hurts: ", {"bold": True, "color": RED}),
     ("17/20 worse, +0.038 CER — thresholding erases faint strokes.")],
    [("Full pipeline worst overall: ", {"bold": True, "color": RED}),
     ("0.359 — stacked neutral stages compound into harm.")],
    ["Denoise / enhance / deskew: neutral within noise."],
    [("Honest bounds: ", {"bold": True}),
     ("n=20, descriptive only, possible train overlap, clean pages only.")],
])
footer(s, 5)

# ---- 6. Closing -----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
title_bar(s, "05 · TAKEAWAY", "Default to the original. Verify everything.")
bullets(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(3.2), [
    "LipiLens works end-to-end and its study answers the question: preprocessing that destroys gray-level stroke detail hurts this VLM.",
    [("Human verification is not a feature — ", {"bold": True}),
     ("it is the architecture (see: the fluent hallucination on MT-003).")],
    "Next: degraded-manuscript robustness, larger ground-truth set, local-GPU path.",
])
text(s, Inches(0.7), Inches(5.4), Inches(11.9), Inches(0.6),
     "Thank you.  Demo: upload → verify → library  ·  Evidence: showcase/GROUND_TRUTH.md",
     size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
footer(s, 6)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")
